"""把挑战A 两轮测试的证据 JSON 转成《测试用例》工作簿。

数据来源（权威运行）：
  _runtime/evidence/challengeA-20261028-194546.json  —— 39 条功能用例
  _runtime/evidence/edge-20261028-194548.json       —— 20 条边界对抗用例
"""
import json
import os
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
EVID = os.path.join(ROOT, "_runtime", "evidence")
OUT = os.path.join(ROOT, "第2次评审测试", "09-挑战A新增功能测试用例.xlsx")

F1 = os.path.join(EVID, "challengeA-20261028-194546.json")
F2 = os.path.join(EVID, "edge-20261028-194548.json")

# ---------------------------------------------------------------- 需求追溯映射
TRACE = {
    "A": "C-29 / NFR-003 / FR-017",
    "B01": "C-06 / FR-017", "B02": "FR-017", "B03": "FR-017",
    "B04": "FR-018", "B05": "FR-018", "B06": "C-07 / FR-019",
    "B07": "C-07 / FR-020", "B08": "FR-018",
    "C": "C-01 / FR-001",
    "D": "C-03 / C-18 / C-27 / FR-005",
    "E01": "C-02 / C-22 / FR-002", "E02": "C-22 / FR-003", "E03": "C-22 / FR-004",
    "F01": "C-12 / FR-022 / FR-023", "F02": "C-12 / NFR-006",
    "F03": "C-12 / FR-022", "F04": "C-17 / FR-027", "F05": "C-26 / FR-030",
    "F06": "C-11 / FR-030",
    "G01": "C-19 / C-28 / FR-013 / FR-028", "G02": "C-19 / FR-034",
    "H": "FR-002（价格校验，回归 BUG-001）",
    "I01": "FR-021", "I02": "FR-021 / C-07", "I03": "NFR-004 / FR-025",
    "J01": "C-03 / C-27", "J02": "C-19 / FR-014",
    "K": "C-29 / NFR-003 / FR-037",
    "L01": "C-13 / FR-026", "L02": "C-13 / FR-026",
    "M": "接口设计 §1（分页参数）",
}

PRECOND = {
    "密码强度": "后端已启动；数据库为空",
    "客户账号": "后端已启动",
    "访问控制": "后端已启动；存在在售商品",
    "多商品": "已登录卖家；已建二级类别",
    "库存": "已登录卖家；已建二级类别",
    "类别": "已登录卖家；已建一级/二级类别",
    "队列": "已登录卖家 + 两个客户；存在在售商品",
    "交易": "商品队列中已有排队意向",
    "售罄": "商品库存为 1，有两名客户排队",
    "价格边界": "已登录卖家；已建二级类别",
    "意向冲突": "已登录客户；存在在售商品",
    "越权": "两个客户各有一条意向",
    "卖家改密": "已用 seller/seller123 登录",
    "冻结": "已登录卖家；商品在售",
    "分页": "后端已启动",
}


def trace_of(cid):
    if cid in TRACE:
        return TRACE[cid]
    return TRACE.get(cid[:1], "")


# ---------------------------------------------------------------- 读证据
def load(path):
    with open(path, encoding="utf-8") as fh:
        return json.load(fh)


d1 = load(F1)
d2 = load(F2)
rows = d1["results"] + d2["results"]

# ---------------------------------------------------------------- 样式
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="D9E2F3")
HEAD_FONT = Font(bold=True, size=10)
BODY_FONT = Font(size=9)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
PASS_FILL = PatternFill("solid", fgColor="E2EFDA")
FAIL_FILL = PatternFill("solid", fgColor="FCE4E4")


def style_header(ws, row=1):
    for c in ws[row]:
        if c.value is None:
            continue
        c.fill = HEAD_FILL
        c.font = HEAD_FONT
        c.alignment = CENTER
        c.border = BORDER


def style_body(ws, start=2):
    for r in ws.iter_rows(min_row=start):
        for c in r:
            if c.value is None:
                continue
            c.font = BODY_FONT
            c.border = BORDER
            c.alignment = WRAP


def widths(ws, ws_w):
    for i, w in enumerate(ws_w, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


wb = Workbook()

# ---------------------------------------------------------------- 表1 用例总览
ws1 = wb.active
ws1.title = "用例总览"
ws1.append(["《在线购物系统（单卖家版）》挑战A 新增功能测试用例总览"])
ws1.append(["被测对象：backend-java（Spring Boot 2.7.18，30 个 REST 接口）　代码版本：main = 8d10723"])
ws1.append(["测试方式：对运行中的后端（127.0.0.1:8000）发真实 HTTP 请求　执行日期：2026-10-28"])
ws1.append(["编写人：缪茂锦（测试岗）　证据文件：challengeA-20261028-194546.json / edge-20261028-194548.json"])
ws1.append([])
ws1.append(["序号", "测试模块", "用例数", "通过", "失败", "通过率", "需求依据"])
style_header(ws1, row=6)

order = []
seen = {}
for r in rows:
    m = r["module"]
    if m not in seen:
        seen[m] = {"n": 0, "p": 0, "fail": 0}
        order.append(m)
    seen[m]["n"] += 1
    if r["result"] == "PASS":
        seen[m]["p"] += 1
    else:
        seen[m]["fail"] += 1

BASIS = {
    "密码强度": "C-29 / NFR-003", "客户账号": "C-06 / FR-017~018",
    "访问控制": "C-07 / FR-019~020", "多商品": "C-01", "库存": "C-03 / C-18 / C-27",
    "类别": "C-02 / C-22", "队列": "C-12 / FR-022", "交易": "C-17 / C-26",
    "售罄": "C-19", "价格边界": "FR-002（回归 BUG-001）", "意向冲突": "FR-021",
    "越权": "NFR-004", "卖家改密": "C-29 / FR-037", "冻结": "C-13", "分页": "接口设计§1",
}
for i, m in enumerate(order, 1):
    s = seen[m]
    ws1.append([i, m, s["n"], s["p"], s["fail"],
                "%.1f%%" % (100.0 * s["p"] / s["n"]), BASIS.get(m, "")])

total = len(rows)
passed = sum(1 for r in rows if r["result"] == "PASS")
ws1.append(["合计", "", total, passed, total - passed,
            "%.1f%%" % (100.0 * passed / total), "两轮合计"])
style_body(ws1, start=7)
for c in ws1[1]:
    c.font = Font(bold=True, size=12)
widths(ws1, [8, 16, 10, 8, 8, 10, 34])

# 结论区
ws1.append([])
ws1.append(["测试结论"])
ws1.cell(row=ws1.max_row, column=1).font = Font(bold=True, size=11)
for line in [
    "1. 挑战A 新增功能在接口层面实现正确：密码强度、客户账号、访问控制、多商品、库存、类别、队列、交易、售罄全部通过。",
    "2. 旧缺陷 BUG-001 已修复：Java 版对 price=0.001 返回 422 而非 500（用例 H02）。",
    "3. 工程层面另有 5 项问题（团队集成测试编译失败、需求编号重复、缺草稿状态、种子密码不可回退、schema 缺 CHECK）。",
    "4. 全部用例均可由第三方用同样的请求体重放复核。",
]:
    ws1.append([line])

# ---------------------------------------------------------------- 表2 测试用例明细
ws2 = wb.create_sheet("测试用例明细")
ws2.append(["用例编号", "所属模块", "用例标题", "需求追溯", "前置条件",
            "测试步骤与数据", "预期结果", "实际结果", "执行结果", "关联缺陷", "备注"])
for r in rows:
    req = r["request"]
    body = req.get("body")
    step = "%s %s" % (req["method"], req["path"])
    if body is not None:
        step += "\n" + json.dumps(body, ensure_ascii=False)
    ws2.append([
        r["id"], r["module"], r["title"], trace_of(r["id"]),
        PRECOND.get(r["module"], ""), step,
        r["expected"],
        "HTTP %s　%s" % (r["http_status"], r["actual"]),
        r["result"], "", r["time"],
    ])
style_header(ws2)
style_body(ws2)
for i, r in enumerate(rows, start=2):
    ws2.cell(row=i, column=9).fill = PASS_FILL if r["result"] == "PASS" else FAIL_FILL
widths(ws2, [10, 14, 34, 26, 26, 46, 30, 34, 12, 10, 24])
ws2.freeze_panes = "A2"

# ---------------------------------------------------------------- 表3 原始响应（证据）
ws3 = wb.create_sheet("原始响应证据")
ws3.append(["用例编号", "请求时间", "HTTP", "完整响应体（供第三方复核）"])
for r in rows:
    resp = r.get("response")
    txt = json.dumps(resp, ensure_ascii=False)
    if len(txt) > 1200:
        txt = txt[:1200] + " …（已截断）"
    ws3.append([r["id"], r["time"], r["http_status"], txt])
style_header(ws3)
style_body(ws3)
widths(ws3, [10, 24, 8, 130])
ws3.freeze_panes = "A2"

# ---------------------------------------------------------------- 表4 环境与复核
ws4 = wb.create_sheet("环境与复核方法")
ws4.append(["项目", "内容"])
info = [
    ("被测系统", "在线购物系统（单卖家版）挑战A 版"),
    ("后端", "backend-java（Java / Spring Boot 2.7.18 / SQLite）"),
    ("构建命令", "mvn -s _runtime/maven-settings.xml \"-Dmaven.test.skip=true\" package"),
    ("启动命令", "SHOP_DATA_DIR=<空目录> java -jar target/shop-backend-1.0.0.jar --server.port=8000"),
    ("服务地址", "http://127.0.0.1:8000　（接口文档 /docs）"),
    ("运行环境", "JDK 25.0.1 + Maven 3.9.6 + Windows"),
    ("测试脚本", "_runtime/challengeA_suite.py（39 条功能）、_runtime/challengeA_edge.py（20 条边界）"),
    ("证据文件", "challengeA-20261028-194546.json / .csv　edge-20261028-194548.json"),
    ("数据隔离", "每轮测试前清空 SHOP_DATA_DIR；测试数据带时间戳后缀，保证可重复"),
    ("", ""),
    ("复核方法一", "打开本工作簿『原始响应证据』页，任选一条，用同样请求体重放，比对返回"),
    ("复核方法二", "按上面的启动命令起后端，运行两个测试脚本，应复现同样结果"),
    ("复核方法三", "用 Swagger UI（http://127.0.0.1:8000/docs）手工发单条请求"),
    ("", ""),
    ("预期值的来源", "来自需求条款（C-xx / FR-xxx / NFR-xxx），不来自代码实现"),
    ("为什么重要", "若先读代码再抄成预期，测试退化为复述实现，永远 100% 通过，无验证价值"),
]
for k, v in info:
    ws4.append([k, v])
style_header(ws4)
style_body(ws4)
widths(ws4, [18, 100])
for i in range(2, ws4.max_row + 1):
    v = ws4.cell(row=i, column=1).value
    if v and str(v).startswith("复核") or v in ("预期值的来源", "为什么重要"):
        ws4.cell(row=i, column=1).font = Font(bold=True, size=9)

os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print("已生成:", os.path.relpath(OUT, ROOT))
print("用例总数:", total, " 通过:", passed, " 失败:", total - passed)
print("模块数:", len(order), "->", order)
