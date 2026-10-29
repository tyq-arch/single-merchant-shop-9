"""把审计结果写成 UTF-8 文本，便于查看。"""
import collections
import os

from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
wb = load_workbook(os.path.join(ROOT, "第2次评审测试", "测试材料", "04-测试过程审计.xlsx"))

out = []
ws = wb["逐条审计"]
rows = list(ws.iter_rows(min_row=2, values_only=True))
out.append("审计条数: %d" % len(rows))
out.append("")
out.append("=== 按假通过风险 ===")
for k, v in collections.Counter(str(r[9]) for r in rows).items():
    out.append("  %-18s %d 条" % (k, v))
out.append("")
out.append("=== 按断言强度 ===")
for k, v in collections.Counter(str(r[7]) for r in rows).items():
    out.append("  %-24s %d 条" % (k, v))
out.append("")
out.append("=== 假通过风险清单 ===")
for r in wb["假通过风险清单"].iter_rows(values_only=True):
    if r[0]:
        out.append("  %s | %s 条 | %s" % (r[0], r[1], str(r[2])[:80]))
out.append("")
out.append("=== 需要人工看的条目 ===")
for r in rows:
    if str(r[9]).startswith("❌") or str(r[9]).startswith("⚠"):
        out.append("  %-8s %-14s %s" % (r[0], str(r[1])[:12], str(r[9])))
out.append("")
out.append("=== 弱断言但复验通过的样例（前 8 条）===")
n = 0
for r in rows:
    if str(r[7]).startswith("弱") and str(r[9]) == "无":
        out.append("  %-6s 预期=%s" % (r[0], str(r[5])[:40]))
        out.append("         实际=%s" % str(r[6])[:130].replace("\n", " "))
        n += 1
        if n >= 8:
            break

txt = "\n".join(out)
open(os.path.join(ROOT, "_runtime", "audit-summary.txt"), "w", encoding="utf-8").write(txt)
print("done")
